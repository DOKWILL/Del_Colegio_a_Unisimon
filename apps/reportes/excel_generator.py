"""
Generador de reportes Excel con openpyxl.

Genera reportes académicos en formato Excel con certificación y nivelación.
"""
import io
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill


def generar_reporte_excel(asignacion, notas_data, asistencia_data, buffer=None):
    """
    Genera un reporte Excel profesional para una asignación.

    Args:
        asignacion: Objeto Asignacion
        notas_data: Lista de dicts con datos de cada estudiante
        asistencia_data: Dict con estadísticas
        buffer: BytesIO buffer

    Returns:
        BytesIO buffer con el Excel generado
    """
    if buffer is None:
        buffer = io.BytesIO()

    wb = Workbook()
    ws = wb.active
    ws.title = 'Reporte Académico'

    # Estilos
    header_font = Font(name='Calibri', bold=True, size=14, color='09843B')
    subtitle_font = Font(name='Calibri', bold=True, size=11, color='23987E')
    info_font = Font(name='Calibri', size=10)
    info_bold = Font(name='Calibri', size=10, bold=True)
    table_header_font = Font(name='Calibri', bold=True, size=9, color='FFFFFF')
    table_header_fill = PatternFill(start_color='09843B', end_color='09843B', fill_type='solid')
    alt_row_fill = PatternFill(start_color='F8F9FA', end_color='F8F9FA', fill_type='solid')
    aprobado_font = Font(name='Calibri', size=10, color='28A745', bold=True)
    reprobado_font = Font(name='Calibri', size=10, color='DC3545', bold=True)
    cert_homol_font = Font(name='Calibri', size=9, color='155724', bold=True)
    cert_homol_fill = PatternFill(start_color='D4EDDA', end_color='D4EDDA', fill_type='solid')
    cert_font = Font(name='Calibri', size=9, color='0C5460', bold=True)
    cert_fill = PatternFill(start_color='D1ECF1', end_color='D1ECF1', fill_type='solid')
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)
    left_wrap = Alignment(horizontal='left', vertical='center', wrap_text=True)
    thin_border = Border(
        left=Side(style='thin', color='CCCCCC'),
        right=Side(style='thin', color='CCCCCC'),
        top=Side(style='thin', color='CCCCCC'),
        bottom=Side(style='thin', color='CCCCCC'),
    )

    # === ENCABEZADO ===
    ws.merge_cells('A1:J1')
    ws['A1'] = 'UNIVERSIDAD SIMÓN BOLÍVAR'
    ws['A1'].font = header_font
    ws['A1'].alignment = center

    ws.merge_cells('A2:J2')
    ws['A2'] = 'DEL COLEGIO A UNISIMÓN — Reporte de Calificaciones, Asistencia y Certificación'
    ws['A2'].font = subtitle_font
    ws['A2'].alignment = center

    # === INFORMACIÓN ===
    row = 4
    info_fields = [
        ('Programa:', asignacion.programa.nombre, 'Periodo:', asignacion.periodo),
        ('Materia:', asignacion.materia.nombre, 'Código:', asignacion.materia.codigo),
        ('Profesor:', asignacion.profesor.nombres, 'Fecha:', date.today().strftime('%d/%m/%Y')),
    ]
    for fields in info_fields:
        ws.cell(row=row, column=1, value=fields[0]).font = info_bold
        ws.cell(row=row, column=2, value=fields[1]).font = info_font
        ws.cell(row=row, column=6, value=fields[2]).font = info_bold
        ws.cell(row=row, column=7, value=fields[3]).font = info_font
        row += 1

    # === TABLA DE CALIFICACIONES ===
    row += 1
    headers = ['#', 'Identificación', 'Estudiante', 'Parcial 1 (30%)',
               'Parcial 2 (30%)', 'Parcial 3 (40%)', 'Definitiva',
               '% Asistencia', 'Estado', 'Certificación']

    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = table_header_font
        cell.fill = table_header_fill
        cell.alignment = center
        cell.border = thin_border

    for i, dato in enumerate(notas_data, 1):
        row += 1
        asist_text = f"{dato['asistencia_pct']:.1f}%"
        if dato.get('aplica_nivelacion'):
            asist_text += ' *'

        values = [
            i,
            dato['identificacion'],
            dato['nombre'],
            dato['parcial1'],
            dato['parcial2'],
            dato['parcial3'],
            dato['definitiva'],
            asist_text,
            dato['estado'] or '-',
            dato.get('certificacion', '') or '-',
        ]
        for col, val in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=val)
            cell.font = info_font
            cell.border = thin_border
            if col >= 3:
                cell.alignment = center
            if col == 3:
                cell.alignment = left_wrap
            if i % 2 == 0:
                cell.fill = alt_row_fill

        # Colorear estado
        estado_cell = ws.cell(row=row, column=9)
        if dato['estado'] == 'Aprobado':
            estado_cell.font = aprobado_font
        elif dato['estado'] == 'Reprobado':
            estado_cell.font = reprobado_font

        # Colorear certificación
        cert_cell = ws.cell(row=row, column=10)
        cert_val = dato.get('certificacion', '')
        if cert_val == 'Certificable y Homologable':
            cert_cell.font = cert_homol_font
            cert_cell.fill = cert_homol_fill
        elif cert_val == 'Certificable':
            cert_cell.font = cert_font
            cert_cell.fill = cert_fill

    # === LEYENDA ===
    row += 2
    ws.cell(row=row, column=1, value='* Porcentaje ajustado por encuentro de nivelación').font = Font(
        name='Calibri', size=9, italic=True, color='666666')
    row += 1
    ws.cell(row=row, column=1, value='7 encuentros regulares + 1 encuentro de nivelación').font = Font(
        name='Calibri', size=9, italic=True, color='666666')

    # === RESUMEN ===
    row += 2
    ws.cell(row=row, column=1, value='RESUMEN DEL CURSO').font = subtitle_font
    row += 1
    resumen_fields = [
        ('Cantidad de estudiantes:', asistencia_data.get('total_estudiantes', 0)),
        ('Aprobados:', asistencia_data.get('aprobados', 0)),
        ('Reprobados:', asistencia_data.get('reprobados', 0)),
        ('Promedio del curso:', f"{asistencia_data.get('promedio_curso', 0):.2f}"),
        ('Promedio de asistencia:', f"{asistencia_data.get('promedio_asistencia', 0):.1f}%"),
    ]
    for label, value in resumen_fields:
        ws.cell(row=row, column=1, value=label).font = info_bold
        ws.cell(row=row, column=2, value=value).font = info_font
        row += 1

    # Ajustar anchos de columna
    col_widths = [5, 16, 30, 14, 14, 14, 12, 14, 12, 28]
    for i, width in enumerate(col_widths, 1):
        ws.column_dimensions[chr(64 + i)].width = width

    wb.save(buffer)
    buffer.seek(0)
    return buffer
